from datetime import datetime, timedelta
from types import SimpleNamespace

import pandas as pd
import pytest

from cache import MemoryCache, cached
from journal_data import load_journal_table, source_digest
from workflow_exports import publish_review
from workflow_runtime import IntelligentCache, content_key


def test_none_result_and_keyword_order_are_cached():
    calls = []

    @cached(MemoryCache())
    def operation(**kwargs):
        calls.append(kwargs)
        return None

    assert operation(first=1, second=2) is None
    assert operation(second=2, first=1) is None
    assert len(calls) == 1


def test_zero_capacity_and_negative_capacity():
    cache = MemoryCache(maxsize=0)
    cache.set('a', 1)
    assert cache.size() == 0
    with pytest.raises(ValueError):
        MemoryCache(maxsize=-1)


def test_content_cache_keys_include_evidence_not_only_count():
    assert content_key('outline', [{'pmid': '1'}]) != content_key('outline', [{'pmid': '2'}])
    assert content_key('intent', {'a': 1, 'b': 2}) == content_key('intent', {'b': 2, 'a': 1})


def test_ai_response_ttl(tmp_path):
    cache = IntelligentCache(str(tmp_path))
    cache.cache_ai_response('old', 'value')
    cache.ai_response_cache['old']['timestamp'] = (datetime.now() - timedelta(hours=2)).isoformat()
    assert cache.get_cached_ai_response('old') is None


def test_journal_loader_reads_all_chunks_and_hash_changes(tmp_path):
    path = tmp_path / 'journal.csv'
    pd.DataFrame({'ISSN': [f'{index:08}' for index in range(2501)]}).to_csv(path, index=False)
    frame = load_journal_table(path, lambda chunk: chunk)
    assert len(frame) == 2501
    previous = source_digest([path])
    with path.open('a') as stream:
        stream.write('99999999\n')
    assert source_digest([path]) != previous


@pytest.mark.parametrize('format_name', ['md', 'docx', 'both'])
def test_cached_review_exports_requested_formats(tmp_path, format_name):
    from docx import Document

    def save_article(content, filename, query, export_docx, export_md):
        md_path = tmp_path / 'review.md' if export_md else None
        docx_path = tmp_path / 'review.docx' if export_docx else None
        if md_path:
            md_path.write_text(content, encoding='utf-8')
        if docx_path:
            document = Document()
            document.add_paragraph(content)
            document.save(docx_path)
        return md_path, docx_path

    paths = publish_review(SimpleNamespace(save_article=save_article), 'o', 'l', 'title', 'review.md', 'q', format_name, cached_content='offline evidence')
    assert bool(paths[0]) == (format_name in {'md', 'both'})
    assert bool(paths[1]) == (format_name in {'docx', 'both'})
    if paths[1]:
        assert Document(paths[1]).paragraphs[0].text == '# title\n\noffline evidence'


def test_export_rejects_nonexistent_files():
    generator = SimpleNamespace(generate_from_files=lambda **kwargs: ('missing.md', None))
    with pytest.raises(RuntimeError, match='未找到'):
        publish_review(generator, 'o', 'l', 't', 'r', 'q', 'md')


def test_real_article_save_contract(tmp_path):
    from medical_review_generator import MedicalReviewGenerator
    generator = MedicalReviewGenerator.__new__(MedicalReviewGenerator)
    generator.output_dir = str(tmp_path)
    paths = publish_review(generator, 'o', 'l', 't', 'review.md', 'q', 'md', cached_content='真实保存路径测试')
    assert paths[0] == str(tmp_path / 'review.md')
    assert (tmp_path / 'review.md').read_text(encoding='utf-8') == '# t\n\n真实保存路径测试'


def test_constructor_preserves_state_file(tmp_path, monkeypatch):
    from intelligent_literature_system import IntelligentLiteratureSystem
    monkeypatch.chdir(tmp_path)
    path = tmp_path / 'system_state.json'
    path.write_text('{"processing": true}', encoding='utf-8')
    IntelligentLiteratureSystem(interactive_mode=False)
    assert path.read_text(encoding='utf-8') == '{"processing": true}'
